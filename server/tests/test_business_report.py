"""The business report: window maths, the policy detail, and that both
renderers actually produce a file.

The report replaced the "month pack" on 2026-08-04. What changed is mostly the
POLICY detail — the pack's policy sheet carried no reward rate, no reward base
and no channel-partner name, which are the three columns anybody actually
reconciles a commission with. Most of what is pinned here is that those columns
exist, mean what they say, and read the same in the Excel and the PDF.

No DB — the renderers take a plain dict from `collect()`, so they are fed stubs.
That is the point of splitting collect() from build_*(): the formatting can be
tested without a database, and the reads can be changed without touching the
layout.
"""

import io
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace

from openpyxl import load_workbook

from app.core.enums import LedgerTxnType, PartyType, PolicyStatus, RewardBasis
from app.services import business_report as report
from app.services import business_report_pdf
from app.services.finance_reports import IST


# --- Window maths ------------------------------------------------------------------


def test_a_month_runs_from_the_1st_to_the_last_instant_in_ist():
    lo, hi, label = report.month_window("2026-07")
    assert label == "July 2026"
    assert (lo.year, lo.month, lo.day, lo.hour) == (2026, 7, 1, 0)
    assert lo.utcoffset() == timedelta(hours=5, minutes=30)
    # Inclusive end: 23:59:59 on the 31st, not midnight on the 1st of August
    # (which would pull the next month's first second into this report).
    assert (hi.year, hi.month, hi.day) == (2026, 7, 31)
    assert (hi.hour, hi.minute, hi.second) == (23, 59, 59)


def test_december_rolls_into_the_next_year():
    lo, hi, label = report.month_window("2026-12")
    assert label == "December 2026"
    assert (hi.year, hi.month, hi.day) == (2026, 12, 31)


def test_february_knows_about_leap_years():
    _lo, hi, _label = report.month_window("2028-02")
    assert hi.day == 29


def test_the_period_presets_are_the_apps_own():
    """Not a second copy. `resolve_period` is the one place that knows what
    "last 3 months" means and that a year here is the Indian financial year —
    a private copy would eventually disagree with the charts sitting directly
    above the download button."""
    now = datetime(2026, 8, 4, 12, 0, tzinfo=IST)
    lo, _hi, label = report.report_window("last_month", now)
    assert (lo.year, lo.month, lo.day) == (2026, 7, 1)
    assert label == "Last month"

    lo3, _hi3, label3 = report.report_window("last_3_months", now)
    assert (lo3.year, lo3.month, lo3.day) == (2026, 6, 1)
    assert label3 == "Last 3 months"

    # "This year" is the FY that started 1 April, not 1 January.
    fy_lo, _fy_hi, fy_label = report.report_window("this_year", now)
    assert (fy_lo.month, fy_lo.day) == (4, 1)
    assert fy_label.startswith("FY")


def test_a_custom_range_is_taken_as_given():
    now = datetime(2026, 8, 4, 12, 0, tzinfo=IST)
    lo, hi, _label = report.report_window(
        "custom", now,
        date_from=datetime(2026, 5, 10, tzinfo=IST),
        date_to=datetime(2026, 6, 20, 23, 59, 59, tzinfo=IST))
    assert (lo.month, lo.day) == (5, 10)
    assert (hi.month, hi.day) == (6, 20)


# --- The router's half of the window ---------------------------------------------------


def test_a_picked_date_is_read_as_an_INDIAN_wall_clock_time():
    """THE off-by-five-and-a-half-hours. `_parse_dt` attaches UTC to a value
    with no offset, which is right for an API timestamp and wrong for a date
    somebody picked out of a calendar: "2026-07-01" would start the range at
    05:30 IST and silently drop everything booked before breakfast."""
    from app.routers.reports import _parse_ist

    got = _parse_ist("2026-07-01")
    assert got.utcoffset() == timedelta(hours=5, minutes=30)
    assert (got.year, got.month, got.day, got.hour) == (2026, 7, 1, 0)


def test_an_explicit_offset_is_still_honoured():
    from app.routers.reports import _parse_ist

    got = _parse_ist("2026-07-01T00:00:00+00:00")
    assert got.utcoffset() == timedelta(0)


def test_a_range_ending_on_a_date_covers_the_whole_of_that_day():
    """A range ending "31 Jul" that stopped at midnight would drop the last
    day's business — which is most of a month-end."""
    from app.routers.reports import _report_window

    _lo, hi, _label = _report_window("custom", "2026-07-01", "2026-07-31")
    assert (hi.day, hi.hour, hi.minute, hi.second) == (31, 23, 59, 59)


def test_a_backwards_range_is_refused():
    from fastapi import HTTPException

    from app.routers.reports import _report_window

    try:
        _report_window("custom", "2026-07-31", "2026-07-01")
    except HTTPException as e:
        assert e.status_code == 400
    else:
        raise AssertionError("a backwards range should be refused")


def test_an_unknown_period_is_refused_by_name():
    from fastapi import HTTPException

    from app.routers.reports import _report_window

    try:
        _report_window("last_fortnight", None, None)
    except HTTPException as e:
        assert e.status_code == 400
        assert "last_3_months" in e.detail      # says what IS allowed
    else:
        raise AssertionError("an unknown period should be refused")


# --- Money and file naming ------------------------------------------------------------


def test_money_is_exported_as_a_number_so_excel_can_sum_it():
    """A formatted string looks right and cannot be added up."""
    assert report.rupees(123456) == 1234.56
    assert isinstance(report.rupees(100), float)
    assert report.rupees(None) == 0


def test_a_whole_month_is_named_by_the_month():
    lo, hi, _l = report.month_window("2026-07")
    assert report.filename_stem(lo, hi) == "Agastya-Report-Jul-2026"


def test_a_range_spells_out_both_ends():
    """`report.xlsx` is unfindable a month later in a folder of twenty."""
    lo = datetime(2026, 7, 1, tzinfo=IST)
    hi = datetime(2026, 9, 30, 23, 59, 59, tzinfo=IST)
    assert report.filename_stem(lo, hi) == "Agastya-Report-01Jul2026-30Sep2026"


def test_a_part_month_is_not_named_as_the_whole_month():
    """1-15 July must not download as "Jul-2026" and get filed as the month."""
    lo = datetime(2026, 7, 1, tzinfo=IST)
    hi = datetime(2026, 7, 15, 23, 59, 59, tzinfo=IST)
    assert report.filename_stem(lo, hi) == "Agastya-Report-01Jul2026-15Jul2026"


# --- Fixtures -------------------------------------------------------------------------


def _spec(key, label, type_="text"):
    return SimpleNamespace(key=key, label=label, type=type_)


def _stub_data(*, can_view_profit=True):
    lo, hi, _label = report.month_window("2026-07")
    policy = SimpleNamespace(
        id="p1", code="AG-POL-000001", policy_number="POL/1",
        created_at=lo + timedelta(days=2), start_date=lo + timedelta(days=2),
        expiry_date=hi + timedelta(days=300), customer_id="c1",
        category_key="motor", subcategory_path=["four_wheeler"],
        insurer_id="i1", broker_id="b1",
        owner_user_id="u1", partner_id="pt1", manager_name="Rahul S",
        status=PolicyStatus.ACTIVE, premium_amount=1_000_000,
        commissionable_premium=847_500, reward_base_field="od_amount",
        details={"vehicle_no": "RJ27BU1692", "od_amount": 500_000},
        renewed_from_policy_id=None,
        reward=SimpleNamespace(
            agency_basis=RewardBasis.PERCENT, agency_value=4000,
            partner_basis=RewardBasis.PERCENT, partner_value=2500))
    finance = SimpleNamespace(
        policy_id="p1", agency_reward=200_000, partner_share=125_000,
        commissionable=500_000, discount=0,
        house_profit=75_000 if can_view_profit else 0, gross_premium=1_000_000)
    txn = SimpleNamespace(
        occurred_at=lo + timedelta(days=3),
        txn_type=LedgerTxnType.PREMIUM_COLLECTED,
        party_type=PartyType.CUSTOMER, party_id="c1", amount_paise=-1_000_000,
        bank_account_id="a1", bank_delta_paise=1_000_000, reference="UTR1",
        note="", policy_id="p1", expense_category=None, paid_to_name=None,
        paid_to_employee_id=None)
    expense = SimpleNamespace(
        occurred_at=lo + timedelta(days=4), txn_type=LedgerTxnType.EXPENSE,
        party_type=PartyType.EXPENSE, party_id="house", amount_paise=-50_000,
        bank_account_id="a1", bank_delta_paise=-50_000, reference="",
        note="July rent", policy_id=None, expense_category="rent",
        paid_to_name="Landlord", paid_to_employee_id=None)
    tds = SimpleNamespace(
        occurred_at=lo + timedelta(days=5), broker_id="b1", policy_id="p1",
        gross_reward_paise=200_000, tds_percent=200, tds_paise=4_000,
        reference="TDS1")
    motor = SimpleNamespace(
        key="motor", label="Motor",
        children=[SimpleNamespace(key="four_wheeler", label="Four Wheeler",
                                  children=[], active=True)],
        custom_fields=[_spec("vehicle_no", "Vehicle No"),
                       _spec("od_amount", "OD Amount", "amount")])
    return {
        "lo": lo, "hi": hi,
        "policies": [policy], "finance": {"p1": finance},
        "rewards": {"p1": "pending"},
        "txns": [txn, expense], "tds": [tds],
        "leads": [SimpleNamespace(
            code="AG-LED-1", name="Prospect", mobile="9999999999",
            interested_in="Motor", category_key=None, stage="new",
            created_by_name="Asha", created_at=lo + timedelta(days=1))],
        "customers": {"c1": SimpleNamespace(id="c1", name="Ravi Kumar",
                                            code="AG-CUS-1")},
        "brokers": {"b1": SimpleNamespace(id="b1", name="Acme Brokers")},
        "insurers": {"i1": SimpleNamespace(id="i1", name="HDFC Ergo")},
        "categories": {"motor": "Motor"},
        "category_docs": {"motor": motor},
        "users": {"u1": SimpleNamespace(id="u1", full_name="Asha Rao",
                                        code="AG-EMP-1"),
                  "pt1": SimpleNamespace(id="pt1", full_name="Partner One",
                                         code="AG-CP-1")},
        "accounts": {"a1": SimpleNamespace(id="a1", name="HDFC Current")},
        "bank_accounts": [SimpleNamespace(
            id="a1", name="HDFC Current", account_type="bank",
            bank_name="HDFC", balance_paise=4_613_000, active=True,
            is_cash_asset=True, last_txn_at=lo + timedelta(days=3))],
        "party_accounts": [
            SimpleNamespace(party_type=PartyType.CUSTOMER, party_id="c1",
                            balance_paise=250_000),
            SimpleNamespace(party_type=PartyType.BROKER, party_id="b1",
                            balance_paise=-90_000),
        ],
        "partner_ledgers": [{
            "partner_id": "pt1", "name": "Partner One", "code": "AG-CP-1",
            "opening": 0, "total_credit": 125_000, "total_debit": 40_000,
            "closing": 85_000,
            "rows": [
                {"date": lo + timedelta(days=2), "label": "Reward earned",
                 "reference": None, "policy": "POL/1", "credit": 125_000,
                 "debit": 0, "balance": 125_000},
                {"date": lo + timedelta(days=9), "label": "Reward paid to you",
                 "reference": "UTR9", "policy": None, "credit": 0,
                 "debit": 40_000, "balance": 85_000},
            ],
        }],
        "monthly": report.monthly_rows(
            lo, hi, [policy], {"p1": finance}, [txn, expense]),
        "renewals_due": [],
        "cash": {"premium_collected": 1_000_000, "reward_received": 196_000,
                 "premium_fronted": 900_000, "partner_payouts": 40_000,
                 "refunds": 0, "expenses": 50_000,
                 "net_profit": 108_000 if can_view_profit else 0},
        "can_view_profit": can_view_profit,
    }


def _book(**kw):
    return load_workbook(io.BytesIO(
        report.build_workbook(_stub_data(**kw), "July 2026")))


def _cell(ws, header: str, row: int = 2):
    headers = [c.value for c in ws[1]]
    return ws.cell(row=row, column=headers.index(header) + 1).value


# --- The workbook ---------------------------------------------------------------------


def test_the_workbook_has_every_sheet():
    wb = _book()
    for sheet in ("Summary", "Revenue by month", "Policies", "Motor details",
                  "Transactions", "Expenses", "Rewards & TDS",
                  "Partner payouts", "Partner ledger", "Broker summary",
                  "Insurer summary", "Partner summary", "Customer summary",
                  "Outstanding", "Bank & cash", "Renewals next month",
                  "Leads"):
        assert sheet in wb.sheetnames, sheet


def test_money_is_written_as_numbers_not_strings():
    """A formatted string looks right in the file and cannot be added up — the
    single thing most likely to make an accountant hand it back."""
    value = _cell(_book()["Policies"], "Premium (Rs)")
    assert value == 10000.0
    assert isinstance(value, (int, float))


# --- The policy detail, which is what this pass was for ---------------------------------


def test_a_policy_row_says_what_rate_we_got():
    """The pack had no rate column at all. Without it a reward amount cannot be
    checked against anything."""
    ws = _book()["Policies"]
    assert _cell(ws, "Agency reward %") == "40%"
    assert _cell(ws, "Partner %") == "25%"


def test_a_policy_row_says_what_the_rate_was_APPLIED_to():
    """THE column the owner asked for: "was it the default amount or some
    TP/OD custom amount". A 40% and a 25% are not comparable until you know
    they were percentages of different things."""
    ws = _book()["Policies"]
    assert _cell(ws, "Reward base") == "OD Amount"
    assert _cell(ws, "Reward base amount (Rs)") == 5000.0


def test_the_default_base_says_so_in_words():
    from app.services import policy_columns

    assert policy_columns.reward_base_label(None, None) \
        == "Commissionable (Net Premium)"
    assert policy_columns.reward_base_label("commissionable", None) \
        == "Commissionable (Net Premium)"


def test_a_flat_reward_is_still_shown_as_a_percentage():
    """Owner 2026-07-18: always a %. A column mixing "40%" with "Rs 1,200"
    cannot be read down the page."""
    from app.services import policy_columns

    assert policy_columns.reward_pct(RewardBasis.FLAT, 50_000, 500_000) == "10.00%"


def test_a_flat_reward_on_a_zero_base_does_not_divide_by_zero():
    from app.services import policy_columns

    assert policy_columns.reward_pct(RewardBasis.FLAT, 50_000, 0) == "0%"


def test_a_policy_row_names_the_channel_partner():
    """The pack carried the partner's SHARE but not their NAME, so a payout
    could not be traced back to a person without a second lookup."""
    assert _cell(_book()["Policies"], "Channel partner") == "Partner One"


def test_a_policy_row_carries_the_reward_outcome():
    ws = _book()["Policies"]
    assert _cell(ws, "Reward eligibility") == "Eligible"
    assert _cell(ws, "Reward outcome") == "Pending"


def test_a_not_eligible_policy_says_so():
    from app.services import policy_columns

    assert policy_columns.eligibility("cancelled") == "Not Eligible"
    assert policy_columns.eligibility(None) == "—"


def test_the_policy_sheet_totals_the_money_columns():
    ws = _book()["Policies"]
    last = ws.max_row
    assert _cell(ws, "Premium (Rs)", last) == 10000.0
    assert _cell(ws, "Agency reward (Rs)", last) == 2000.0


def test_the_export_and_the_report_share_one_definition_of_a_rate():
    """They had already drifted once, which is why these live in
    services/policy_columns now."""
    from app.routers import policies as policies_router
    from app.services import policy_columns

    assert policies_router._reward_pct is policy_columns.reward_pct
    assert policies_router._reward_base_label is policy_columns.reward_base_label
    assert policies_router._eligibility is policy_columns.eligibility


# --- Type-specific fields get real columns ----------------------------------------------


def test_a_policy_types_own_fields_become_real_columns():
    """"Vehicle No: RJ27…" squashed into one text cell can be read but not
    sorted, filtered or summed — and an OD column you cannot total is not much
    of a report."""
    ws = _book()["Motor details"]
    assert _cell(ws, "Vehicle No") == "RJ27BU1692"


def test_an_amount_field_comes_out_summable():
    """Stored in paise, printed in rupees AS A NUMBER."""
    value = _cell(_book()["Motor details"], "OD Amount")
    assert value == 5000.0
    assert isinstance(value, (int, float))


def test_a_type_with_no_configured_fields_gets_no_sheet():
    """Otherwise every policy type adds an empty sheet nobody opens."""
    data = _stub_data()
    data["category_docs"]["motor"].custom_fields = []
    wb = load_workbook(io.BytesIO(report.build_workbook(data, "July 2026")))
    assert "Motor details" not in wb.sheetnames


# --- The partner ledger -------------------------------------------------------------------


def test_the_partner_ledger_opens_and_closes():
    ws = _book()["Partner ledger"]
    col = [c.value for c in ws["C"]]
    assert "Opening balance" in col
    assert "Closing balance" in col


def test_the_partner_ledger_runs_a_balance():
    # Excel stores a whole rupee figure as an integer, so both types count.
    ws = _book()["Partner ledger"]
    balances = [c.value for c in ws["G"][1:]
                if isinstance(c.value, (int, float))]
    # opening 0, +1250 earned, -400 paid, closing 850.
    assert balances == [0, 1250, 850, 850]


def test_the_report_and_the_partner_statement_read_one_ledger():
    """A partner holding their statement and the agency holding this report
    disagreeing about the same account is the worst argument to have."""
    import inspect

    from app.services import finance_balance

    assert "partner_ledger(" in inspect.getsource(
        finance_balance.partner_statement)
    assert "finance_balance.partner_ledger" in inspect.getsource(
        report._partner_ledgers)


# --- Revenue by month ---------------------------------------------------------------------


def test_revenue_is_broken_out_by_month():
    """A three-month range reporting only a total hides the thing it was
    chosen to show."""
    lo = datetime(2026, 6, 1, tzinfo=IST)
    hi = datetime(2026, 8, 31, 23, 59, 59, tzinfo=IST)
    policies = [
        SimpleNamespace(id="a", created_at=datetime(2026, 6, 5, tzinfo=IST),
                        premium_amount=100_000),
        SimpleNamespace(id="b", created_at=datetime(2026, 7, 5, tzinfo=IST),
                        premium_amount=300_000),
        SimpleNamespace(id="c", created_at=datetime(2026, 7, 20, tzinfo=IST),
                        premium_amount=200_000),
    ]
    rows = report.monthly_rows(lo, hi, policies, {}, [])
    assert [r["label"] for r in rows] == ["Jun 2026", "Jul 2026"]
    assert [r["policies"] for r in rows] == [1, 2]
    assert rows[1]["premium"] == 500_000


def test_expenses_land_in_the_month_they_were_paid():
    lo = datetime(2026, 6, 1, tzinfo=IST)
    hi = datetime(2026, 7, 31, 23, 59, 59, tzinfo=IST)
    txn = SimpleNamespace(occurred_at=datetime(2026, 7, 9, tzinfo=IST),
                          txn_type=LedgerTxnType.EXPENSE, amount_paise=-25_000)
    rows = report.monthly_rows(lo, hi, [], {}, [txn])
    assert rows[0]["label"] == "Jul 2026"
    assert rows[0]["expenses"] == 25_000


# --- Permission gating ----------------------------------------------------------------------


def test_profit_is_zero_for_a_viewer_without_the_permission():
    assert _cell(_book(can_view_profit=False)["Policies"],
                 "Net profit (Rs)") == 0


def test_the_blanking_happens_once_in_collect_not_per_renderer():
    """Blanking in one renderer and forgetting the other is exactly how a
    permission leaks."""
    import inspect

    src = inspect.getsource(report.collect)
    assert "if not can_view_profit:" in src
    assert "pf.house_profit = 0" in src


# --- What must NOT be in the file --------------------------------------------------------------


def test_no_bank_account_numbers_leave_the_building():
    """This file gets emailed around. The account number is encrypted at rest
    precisely so it does not travel."""
    ws = _book()["Bank & cash"]
    headers = [c.value for c in ws[1]]
    for banned in ("Account number", "IFSC", "UPI"):
        assert banned not in headers


def test_a_leads_phone_number_is_not_in_the_file():
    """The one thing in here worth something to whoever it gets forwarded to."""
    ws = _book()["Leads"]
    assert "Mobile" not in [c.value for c in ws[1]]


# --- Size ceiling -------------------------------------------------------------------------------


def test_the_row_cap_is_a_refusal_not_a_crash():
    """This app runs as ONE small instance; a "Till date" download on a big
    book would exhaust its memory rather than time out."""
    assert report.MAX_POLICIES == 25_000
    assert report.MAX_TRANSACTIONS == 50_000
    assert issubclass(report.ReportTooLarge, Exception)


def test_the_cap_is_checked_before_anything_is_loaded():
    """A limit enforced after the rows are already in memory protects nothing.

    The check runs on `.count()` — which streams nothing — and raises before the
    first `.to_list()`, which is the call that actually hydrates documents.
    """
    import inspect

    src = inspect.getsource(report.collect)
    assert src.index("raise ReportTooLarge") < src.index(".to_list()")


# --- Renderers -----------------------------------------------------------------------------------


def test_the_pdf_builds():
    content = business_report_pdf.build_pdf(_stub_data(), "July 2026")
    assert content[:4] == b"%PDF", "must be a real PDF"
    assert len(content) > 2000


def test_the_pdf_builds_without_profit_too():
    """The permission-gated path renders a different set of tiles AND drops a
    policy column — it has to be exercised, or the first person without the
    flag gets a 500."""
    content = business_report_pdf.build_pdf(
        _stub_data(can_view_profit=False), "July 2026")
    assert content[:4] == b"%PDF"


def test_the_pdf_caps_the_transaction_list():
    """A PDF holding four thousand ledger rows is megabytes to email and
    nobody reads past the second page."""
    assert business_report_pdf.MAX_PDF_TRANSACTIONS == 300
    import inspect

    src = inspect.getsource(business_report_pdf.build_pdf)
    assert "MAX_PDF_TRANSACTIONS" in src
    # …and it says how many it left behind rather than pretending there were
    # no more.
    assert "are in the Excel download" in src


def test_the_pdf_carries_every_policy():
    """The part somebody reads down and ticks off. Leaving it out made the old
    PDF a cover sheet rather than a report."""
    import inspect

    src = inspect.getsource(business_report_pdf.build_pdf)
    assert "policy_rows(data)" in src


def _pdf_pages(content: bytes) -> int:
    import re

    return len(re.findall(rb"/Type\s*/Page[^sC]", content))


def test_the_new_sections_actually_render():
    """A source check proves the CALL is there; this proves it draws something.

    A quiet period is one page — the summary. A period with a policy, a partner
    ledger and transactions opens a page for each, so the count has to grow.
    Without this, a section that silently rendered nothing would still pass
    every other test in this file.
    """
    full = _pdf_pages(business_report_pdf.build_pdf(_stub_data(), "July 2026"))

    empty = _stub_data()
    empty.update(policies=[], finance={}, rewards={}, txns=[], tds=[], leads=[],
                 party_accounts=[], renewals_due=[], partner_ledgers=[],
                 monthly=[], bank_accounts=[],
                 cash={k: 0 for k in empty["cash"]})
    quiet = _pdf_pages(business_report_pdf.build_pdf(empty, "July 2026"))

    assert quiet == 1
    assert full >= quiet + 3, (
        f"expected a page each for policies, partner ledgers and "
        f"transactions; got {full} vs {quiet}")


def test_the_pdf_and_the_excel_print_the_same_policy_rows():
    """One builder, so the reward rate a policy shows cannot differ between the
    two files of the same download."""
    import inspect

    assert "policy_rows" in inspect.getsource(business_report_pdf)
    assert "policy_rows(data)" in inspect.getsource(report.build_workbook)


def test_an_empty_period_still_produces_both_files():
    """A quiet month must download, not explode on an empty table."""
    empty = _stub_data()
    empty.update(policies=[], finance={}, rewards={}, txns=[], tds=[], leads=[],
                 party_accounts=[], renewals_due=[], partner_ledgers=[],
                 monthly=[], bank_accounts=[],
                 cash={k: 0 for k in empty["cash"]})
    assert report.build_workbook(empty, "July 2026")[:2] == b"PK"
    assert business_report_pdf.build_pdf(empty, "July 2026")[:4] == b"%PDF"


def test_the_month_pack_is_gone():
    """Renamed, not kept alongside — two report builders is how the numbers
    start to differ between two files that claim to say the same thing."""
    import importlib

    for dead in ("app.services.month_pack", "app.services.month_pack_pdf"):
        try:
            importlib.import_module(dead)
        except ModuleNotFoundError:
            continue
        raise AssertionError(f"{dead} is still importable")
