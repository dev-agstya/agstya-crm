"""Attendance API — the punch clock, the register, the team board and the
correction queue.

WHOSE ATTENDANCE YOU MAY READ
------------------------------
Every employee always reaches their OWN, with no flag at all — the same rule
targets already follow ("being told your own number is not a privilege", owner
G2). `view_attendance` is what opens somebody ELSE's.

That is enforced by `_target_user`, which is the only way any endpoint in this
file resolves a `user_id`. Reading a request parameter and trusting it is how
"employees can see their own" becomes "employees can see everyone's" by way of a
URL somebody edited.

WHO "THE MANAGER" IS (owner E11)
---------------------------------
There is NO employee-to-employee hierarchy in this app. `reports_to_id` existed
once, was read by nothing, and was deleted on 2026-08-05;
`relationship_manager_id` links a CHANNEL PARTNER to the employee who handles
them and says nothing about staff. The owner chose option (a): "the manager" is
whoever holds `manage_attendance`. No new structure, and the people trusted to
fix a register are exactly the people you would give the flag to.

THE PUNCH IS ALWAYS TODAY AND ALWAYS THE CALLER
------------------------------------------------
No endpoint here takes a date or a user for a punch. A backdated or delegated
self-punch is an attendance system that records nothing (owner C3/C6);
corrections go through the queue, and a manager's edit goes through
`PATCH /day`, which demands a reason and writes an audit row.
"""

from __future__ import annotations

import io
from datetime import datetime
from typing import Optional

from fastapi import (
    APIRouter, Body, Depends, HTTPException, Query, Request, Response, status,
)

from app.core.dependencies import (
    get_active_user, get_inhouse_user, require_permission,
)
from app.core.enums import (
    AttendanceSource, AttendanceStatus, AuditAction, HrRequestStatus,
)
from app.core.permissions import (
    EXPORT_DATA, MANAGE_ATTENDANCE, VIEW_ATTENDANCE, can,
)
from app.models.attendance import AttendanceCorrection, AttendanceDay
from app.models.base import utcnow
from app.models.leave import LeaveRequest
from app.models.user import User
from app.routers._helpers import parse_object_id
from app.schemas.common import Message, Page
from app.schemas.hr import (
    AttendanceDayEdit,
    AttendanceDayOut,
    AttendanceMonthOut,
    BreakOut,
    CorrectionCreate,
    CorrectionOut,
    DecisionIn,
    GeofenceInfo,
    MonthSummary,
    PunchIn,
    PunchState,
    TeamDayOut,
    TeamMemberDay,
    TeamMonthOut,
    TeamMonthRow,
)
from app.services import hr_attendance as svc, hr_calendar as cal
from app.services import hr_geo as geo
from app.services import hr_leave, notifications, settings_svc
from app.services.audit import log_action

# Staff-only. The guard is attached to the WHOLE router, not per endpoint, so a
# new route added here is closed to channel partners by default — see
# core/dependencies.get_inhouse_user. Partners are external and are not covered
# by this module at all (owner A10).
router = APIRouter(prefix="/api/hr/attendance", tags=["attendance"],
                   dependencies=[Depends(get_inhouse_user)])


# --- Access ----------------------------------------------------------------------


async def _target_user(actor: User, user_id: Optional[str]) -> User:
    """The employee an endpoint is about, having checked the caller may see them.

    THE ONE PLACE a `user_id` parameter is turned into a user in this file. Own
    record needs nothing; somebody else's needs `view_attendance`. A missing or
    self-referencing id is the caller.
    """
    if not user_id or user_id == str(actor.id) or user_id == "me":
        return actor
    if not can(actor, VIEW_ATTENDANCE):
        raise HTTPException(
            status.HTTP_403_FORBIDDEN,
            "You can only see your own attendance.")
    target = await User.get(await parse_object_id(user_id))
    if target is None or getattr(target, "is_deleted", False):
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Employee not found.")
    return target


def _month_or_now(month: Optional[str], now: datetime) -> str:
    if month and len(month) == 7 and month[4] == "-":
        try:
            cal.month_bounds(month)
            return month
        except (ValueError, IndexError):
            pass
    return cal.month_key_of(cal.day_key(now))


# --- The punch clock -------------------------------------------------------------


@router.get("/today", response_model=PunchState)
async def my_today(actor: User = Depends(get_active_user)) -> PunchState:
    """The signed-in person's own day. No flag — this is their own record."""
    now = utcnow()
    hr = (await settings_svc.get_settings()).hr
    day_str = cal.day_key(now)
    d = cal.parse_day(day_str)
    policy = svc.policy_for(hr, actor)

    row = await svc.today_for(actor, now=now)
    if row is not None and row.clock_in is not None and row.clock_out is None:
        svc.recompute(row, policy, now=now)

    holidays = await cal.holiday_map(day_str, day_str)
    leave = await LeaveRequest.find_one({
        "user_id": str(actor.id),
        "status": HrRequestStatus.APPROVED.value,
        "start_date": {"$lte": day_str}, "end_date": {"$gte": day_str},
    })
    running = row.open_break if row else None

    return PunchState(
        day=day_str,
        server_now=now,
        clock_in=row.clock_in if row else None,
        clock_out=row.clock_out if row else None,
        on_break=running is not None,
        break_started_at=running.start if running else None,
        breaks=[BreakOut(**b.model_dump()) for b in (row.breaks if row else [])],
        worked_minutes=row.worked_minutes if row else 0,
        break_minutes=row.break_minutes if row else 0,
        status=row.effective_status if row else AttendanceStatus.NOT_MARKED.value,
        is_late=bool(row.is_late) if row else False,
        late_minutes=row.late_minutes if row else 0,
        is_week_off=cal.is_week_off(d, policy.week_off_days),
        holiday_name=holidays.get(day_str),
        on_leave=leave is not None,
        leave_code=leave.code if leave else None,
        shift_start=policy.shift_start,
        shift_end=policy.shift_end,
        can_punch=svc.is_hr_subject(actor),
        work_location=row.work_location if row else None,
        distance_m=(row.punch_in_location.distance_m
                    if row and row.punch_in_location else None),
        # Sent whether or not it is on, because the tile has to know BEFORE the
        # button is pressed whether to ask the browser for a location — asking
        # for it when no geofence exists is a permission prompt with no purpose.
        geofence=GeofenceInfo(
            enabled=geo.is_configured(hr),
            office_label=hr.office_label,
            office_lat=hr.office_lat,
            office_lng=hr.office_lng,
            radius_m=hr.office_radius_m,
            required=bool(hr.geofence_enabled
                          and hr.require_location),
        ),
    )


async def _punch(action, actor: User, request: Request,
                 audit: Optional[AuditAction] = None,
                 summary: str = "",
                 where: Optional[PunchIn] = None) -> PunchState:
    """Shared body of the four punch endpoints.

    Every one of them: refuse a non-employee, run the service, turn its
    ValueError into a sentence, audit, and answer with the SAME `PunchState`
    the GET returns — so the panel never has to guess what changed and never
    needs a second round trip to find out.
    """
    if not svc.is_hr_subject(actor):
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            "Attendance is recorded for employees. Your account does not "
            "clock in.")
    hr = (await settings_svc.get_settings()).hr
    from app.core.rate_limit import client_ip
    # The reported position, forwarded untouched. This router does not classify
    # it — services/hr_geo is the one place that decides what a coordinate
    # means, and a second opinion here is how the punch tile and the month
    # register would start disagreeing about the same morning.
    place = where or PunchIn()
    coords = {"lat": place.lat, "lng": place.lng,
              "accuracy_m": place.accuracy_m}
    try:
        if action is svc.punch_in:
            await action(actor, hr, ip=client_ip(request), **coords)
        elif action is svc.punch_out:
            await action(actor, hr, **coords)
        else:
            # Breaks have no location: they happen inside a day whose place was
            # already decided at clock-in.
            await action(actor, hr)
    except ValueError as exc:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, str(exc))
    if audit is not None:
        await log_action(
            audit, actor_id=str(actor.id), actor_name=actor.full_name,
            actor_role=actor.account_type.value, entity_type="attendance",
            entity_id=str(actor.id), request=request, summary=summary)
    return await my_today(actor)


@router.post("/punch-in", response_model=PunchState)
async def punch_in(request: Request, where: PunchIn = Body(default=PunchIn()),
                   actor: User = Depends(get_active_user)) -> PunchState:
    """Clock in, optionally saying where from.

    The body is OPTIONAL and stays optional. Clients that send no coordinates
    still clock in — the day is simply recorded as unplaceable and flagged for a
    human, which is the honest answer and not the employee's fault. The owner
    can require it (`require_location`) if they would rather refuse the punch.
    """
    return await _punch(svc.punch_in, actor, request,
                        AuditAction.ATTENDANCE_PUNCHED_IN, "Clocked in",
                        where=where)


@router.post("/punch-out", response_model=PunchState)
async def punch_out(request: Request, where: PunchIn = Body(default=PunchIn()),
                    actor: User = Depends(get_active_user)) -> PunchState:
    return await _punch(svc.punch_out, actor, request,
                        AuditAction.ATTENDANCE_PUNCHED_OUT, "Clocked out",
                        where=where)


@router.post("/break/start", response_model=PunchState)
async def break_start(request: Request,
                      actor: User = Depends(get_active_user)) -> PunchState:
    # No audit row: a break is a detail OF the day, it is visible on the day
    # itself, and one row per tea break would drown the trail that matters.
    return await _punch(svc.break_start, actor, request)


@router.post("/break/end", response_model=PunchState)
async def break_end(request: Request,
                    actor: User = Depends(get_active_user)) -> PunchState:
    return await _punch(svc.break_end, actor, request)


# --- The register ----------------------------------------------------------------


def _day_out(v: svc.DayView) -> AttendanceDayOut:
    return AttendanceDayOut(**v.as_dict())


@router.get("/month", response_model=AttendanceMonthOut)
async def month(month: Optional[str] = Query(default=None),
                user_id: Optional[str] = Query(default=None),
                actor: User = Depends(get_active_user)) -> AttendanceMonthOut:
    """One employee's month. Own by default; somebody else's needs the flag."""
    now = utcnow()
    target = await _target_user(actor, user_id)
    hr = (await settings_svc.get_settings()).hr
    key = _month_or_now(month, now)
    days = await svc.build_month(target, key, hr, now=now)
    return AttendanceMonthOut(
        month=key, user_id=str(target.id), user_name=target.full_name,
        days=[_day_out(d) for d in days],
        summary=MonthSummary(**svc.summarise(days, hr)))


@router.get("/team", response_model=TeamDayOut,
            dependencies=[Depends(require_permission(VIEW_ATTENDANCE))])
async def team_day(day: Optional[str] = Query(default=None),
                   actor: User = Depends(get_active_user)) -> TeamDayOut:
    """Today's board: who is in, who is late, who is off, who has not arrived.

    THREE queries for the whole agency, never one per person — the attendance
    rows, the approved leaves and the holiday. The roster itself is the staff
    list (tens of rows), which is the same bounded set `notify_permission`
    already walks; it is not the `find_all()` pattern CLAUDE.md asks new reads
    not to widen.
    """
    now = utcnow()
    hr = (await settings_svc.get_settings()).hr
    key = day if (day and len(day) == 10) else cal.day_key(now)
    try:
        d = cal.parse_day(key)
    except ValueError:
        raise HTTPException(status.HTTP_400_BAD_REQUEST,
                            "That is not a real date.")

    employees = await svc.hr_employees()
    rows = {r.user_id: r for r in
            await AttendanceDay.find({"day": key}).to_list()}
    leaves = {r.user_id: r for r in await LeaveRequest.find({
        "status": HrRequestStatus.APPROVED.value,
        "start_date": {"$lte": key}, "end_date": {"$gte": key},
    }).to_list()}
    holidays = await cal.holiday_map(key, key)
    is_off = cal.is_week_off(d, hr.week_off_days)
    today_key = cal.day_key(now)

    members: list[TeamMemberDay] = []
    for user in employees:
        uid = str(user.id)
        rec = rows.get(uid)
        policy = svc.policy_for(hr, user)
        if rec is not None and rec.clock_in is not None and rec.clock_out is None:
            svc.recompute(rec, policy, now=now)
        view = svc._view_for(d, key, rec, leaves.get(uid), holidays, policy,
                             today=today_key, joined=svc._joining_key(user))
        prof = getattr(user, "employee_profile", None)
        members.append(TeamMemberDay(
            user_id=uid, name=user.full_name, code=user.code,
            designation=getattr(prof, "designation", None),
            status=view.status, clock_in=view.clock_in,
            clock_out=view.clock_out, worked_minutes=view.worked_minutes or 0,
            is_late=bool(view.is_late), late_minutes=view.late_minutes or 0,
            on_break=bool(rec and rec.open_break is not None),
            missed_punch_out=bool(view.missed_punch_out),
            leave_code=view.leave_code))

    # Counts computed HERE, not in the browser. The tiles and the rows under
    # them are the same screen; two derivations of "how many are late" is how
    # they end up saying different things.
    return TeamDayOut(
        day=key, is_week_off=is_off, holiday_name=holidays.get(key),
        members=members,
        in_count=sum(1 for m in members if m.clock_in is not None),
        late_count=sum(1 for m in members if m.is_late),
        leave_count=sum(1 for m in members
                        if m.status in (AttendanceStatus.ON_LEAVE.value,
                                        AttendanceStatus.LEAVE_UNPAID.value)),
        absent_count=sum(1 for m in members
                         if m.status == AttendanceStatus.ABSENT.value),
        not_in_count=sum(1 for m in members if m.clock_in is None
                         and m.status == AttendanceStatus.NOT_MARKED.value))


@router.get("/team/month", response_model=TeamMonthOut,
            dependencies=[Depends(require_permission(VIEW_ATTENDANCE))])
async def team_month(month: Optional[str] = Query(default=None),
                     actor: User = Depends(get_active_user)) -> TeamMonthOut:
    """The month grid: people down the side, days across the top.

    This is the screen the month's pay is worked out from, so it carries every
    person's `MonthSummary` alongside their days — the manager reads
    `payable_days` off the row rather than opening thirty records.
    """
    now = utcnow()
    hr = (await settings_svc.get_settings()).hr
    key = _month_or_now(month, now)
    employees = await svc.hr_employees()

    rows: list[TeamMonthRow] = []
    for user in employees:
        days = await svc.build_month(user, key, hr, now=now)
        prof = getattr(user, "employee_profile", None)
        rows.append(TeamMonthRow(
            user_id=str(user.id), name=user.full_name, code=user.code,
            designation=getattr(prof, "designation", None),
            days=[_day_out(d) for d in days],
            summary=MonthSummary(**svc.summarise(days, hr))))
    return TeamMonthOut(month=key, rows=rows)


# --- A manager's edit ------------------------------------------------------------


@router.patch("/day", response_model=AttendanceDayOut,
              dependencies=[Depends(require_permission(MANAGE_ATTENDANCE))])
async def edit_day(payload: AttendanceDayEdit, request: Request,
                   actor: User = Depends(get_active_user)) -> AttendanceDayOut:
    """Correct one day for one employee (owner C5).

    The reason is mandatory at the schema, and the audit row carries the before
    and after. An attendance edit is one person making a claim about another
    person's hours; at month end it is what the pay is worked out from, so it
    gets the same treatment a ledger correction does.
    """
    target = await User.get(await parse_object_id(payload.user_id))
    if target is None or getattr(target, "is_deleted", False):
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Employee not found.")
    if not svc.is_hr_subject(target):
        raise HTTPException(status.HTTP_400_BAD_REQUEST,
                            "Attendance is only recorded for employees.")

    hr = (await settings_svc.get_settings()).hr
    day = await svc.get_or_create_day(target, payload.day)
    d = cal.parse_day(payload.day)
    before = {"clock_in": day.clock_in, "clock_out": day.clock_out,
              "status": day.effective_status}

    if payload.clock_in is not None:
        day.clock_in = (cal.combine_ist(d, payload.clock_in)
                        if payload.clock_in else None)
    if payload.clock_out is not None:
        day.clock_out = (cal.combine_ist(d, payload.clock_out)
                         if payload.clock_out else None)
    # Setting a clock-out by hand answers the missed-punch question, so the flag
    # comes off — leaving it would keep the day amber in the queue for ever and
    # train people to ignore the colour.
    if payload.clock_out:
        day.missed_punch_out = False
    if payload.status is not None:
        day.manual_status = payload.status or None
    if payload.note is not None:
        day.note = payload.note or None

    day.user_name = target.full_name
    day.source = AttendanceSource.MANUAL.value
    day.edit_reason = payload.reason
    day.edited_by = str(actor.id)
    day.edited_by_name = actor.full_name
    day.edited_at = utcnow()
    svc.recompute(day, svc.policy_for(hr, target))
    await day.save()

    await log_action(
        AuditAction.ATTENDANCE_EDITED, actor_id=str(actor.id),
        actor_name=actor.full_name, actor_role=actor.account_type.value,
        entity_type="attendance", entity_id=str(day.id), request=request,
        summary=f"Edited {target.full_name}'s attendance for {payload.day}",
        meta={"before": {k: str(v) for k, v in before.items()},
              "after": {"clock_in": str(day.clock_in),
                        "clock_out": str(day.clock_out),
                        "status": day.effective_status},
              "reason": payload.reason})

    await notifications.create_notification(
        str(target.id), "Your attendance was updated",
        body=f"{actor.full_name} changed your {payload.day}: "
             f"{payload.reason}",
        category="attendance",
        link=f"/hr/attendance?month={cal.month_key_of(payload.day)}")

    days = await svc.build_month(target, cal.month_key_of(payload.day), hr)
    match = next((v for v in days if v.day == payload.day), None)
    return _day_out(match) if match else AttendanceDayOut(
        day=payload.day, weekday=d.weekday(), status=day.effective_status)


# --- Corrections -----------------------------------------------------------------


@router.get("/corrections", response_model=Page[CorrectionOut])
async def list_corrections(
    mine: bool = Query(default=True),
    status_filter: Optional[str] = Query(default=None, alias="status"),
    page: int = Query(default=1, ge=1),
    page_size: int = Query(default=25, ge=1, le=200),
    actor: User = Depends(get_active_user),
) -> Page[CorrectionOut]:
    """`mine=true` is my own requests; `mine=false` is the approval queue.

    The queue needs `view_attendance`. Defaulting to `mine` means a client that
    forgets the parameter gets the SAFE answer — their own — rather than
    everybody's.

    Paginated (owner 2026-09-12) — a queue running for months easily passes 25
    rows. A plain Mongo query, so count/sort/page all happen at the database.
    """
    query: dict = {}
    if mine or not can(actor, VIEW_ATTENDANCE):
        query["user_id"] = str(actor.id)
    if status_filter:
        query["status"] = status_filter
    total = await AttendanceCorrection.find(query).count()
    rows = await AttendanceCorrection.find(query).sort("-created_at") \
        .skip((page - 1) * page_size).limit(page_size).to_list()
    return Page[CorrectionOut](
        items=[CorrectionOut.from_model(r) for r in rows],
        total=total, page=page, page_size=page_size)


@router.post("/corrections", response_model=CorrectionOut,
             status_code=status.HTTP_201_CREATED)
async def raise_correction(payload: CorrectionCreate, request: Request,
                           actor: User = Depends(get_active_user)
                           ) -> CorrectionOut:
    """Ask for a day to be fixed. Always about the CALLER's own day.

    There is no `user_id` on this schema on purpose: raising a correction for
    somebody else is a manager's edit, which is `PATCH /day` and demands a
    reason on the record.
    """
    if not svc.is_hr_subject(actor):
        raise HTTPException(status.HTTP_400_BAD_REQUEST,
                            "Attendance is only recorded for employees.")
    if payload.day > cal.day_key(utcnow()):
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            "You cannot correct a day that has not happened yet.")

    existing = await AttendanceCorrection.find_one({
        "user_id": str(actor.id), "day": payload.day,
        "status": HrRequestStatus.PENDING.value})
    if existing is not None:
        raise HTTPException(
            status.HTTP_409_CONFLICT,
            f"You already have a correction waiting for {payload.day}.")

    d = cal.parse_day(payload.day)
    row = AttendanceCorrection(
        user_id=str(actor.id), user_name=actor.full_name, day=payload.day,
        requested_clock_in=(cal.combine_ist(d, payload.clock_in)
                            if payload.clock_in else None),
        requested_clock_out=(cal.combine_ist(d, payload.clock_out)
                             if payload.clock_out else None),
        reason=payload.reason)
    await row.insert()

    await log_action(
        AuditAction.ATTENDANCE_CORRECTION_REQUESTED, actor_id=str(actor.id),
        actor_name=actor.full_name, actor_role=actor.account_type.value,
        entity_type="attendance_correction", entity_id=str(row.id),
        request=request,
        summary=f"Requested an attendance correction for {payload.day}")
    await notifications.notify_permission(
        MANAGE_ATTENDANCE, "Attendance correction to review",
        body=f"{actor.full_name} asked for {payload.day} to be corrected.",
        category="attendance", link="/hr/attendance?view=corrections",
        exclude_user_id=str(actor.id))
    return CorrectionOut.from_model(row)


@router.post("/corrections/{correction_id}/decide", response_model=CorrectionOut,
             dependencies=[Depends(require_permission(MANAGE_ATTENDANCE))])
async def decide_correction(correction_id: str, payload: DecisionIn,
                            request: Request,
                            actor: User = Depends(get_active_user)
                            ) -> CorrectionOut:
    """Approve or reject a correction. Approving APPLIES it to the day.

    Applying it here rather than making the manager re-enter the times is the
    entire point of the queue: a request they have to retype by hand is a
    request they will fix directly and leave pending for ever.
    """
    row = await AttendanceCorrection.get(await parse_object_id(correction_id))
    if row is None:
        raise HTTPException(status.HTTP_404_NOT_FOUND,
                            "Correction request not found.")
    if row.status != HrRequestStatus.PENDING.value:
        raise HTTPException(status.HTTP_409_CONFLICT,
                            "That request has already been decided.")
    if row.user_id == str(actor.id):
        raise HTTPException(
            status.HTTP_403_FORBIDDEN,
            "You cannot decide your own correction request.")

    row.status = (HrRequestStatus.APPROVED.value if payload.approve
                  else HrRequestStatus.REJECTED.value)
    row.decided_by = str(actor.id)
    row.decided_by_name = actor.full_name
    row.decided_at = utcnow()
    row.decision_note = payload.note
    row.updated_at = utcnow()
    await row.save()

    if payload.approve:
        target = await User.get(row.user_id)
        if target is not None:
            hr = (await settings_svc.get_settings()).hr
            day = await svc.get_or_create_day(target, row.day)
            if row.requested_clock_in:
                day.clock_in = row.requested_clock_in
            if row.requested_clock_out:
                day.clock_out = row.requested_clock_out
                day.missed_punch_out = False
            day.user_name = target.full_name
            day.source = AttendanceSource.CORRECTION.value
            day.edit_reason = row.reason
            day.edited_by = str(actor.id)
            day.edited_by_name = actor.full_name
            day.edited_at = utcnow()
            svc.recompute(day, svc.policy_for(hr, target))
            await day.save()

    await log_action(
        AuditAction.ATTENDANCE_CORRECTION_DECIDED, actor_id=str(actor.id),
        actor_name=actor.full_name, actor_role=actor.account_type.value,
        entity_type="attendance_correction", entity_id=str(row.id),
        request=request,
        summary=f"{'Approved' if payload.approve else 'Rejected'} "
                f"{row.user_name}'s correction for {row.day}")
    await notifications.create_notification(
        row.user_id,
        f"Attendance correction {'approved' if payload.approve else 'rejected'}",
        body=(payload.note or f"Your {row.day} correction was "
              f"{'applied' if payload.approve else 'not applied'}."),
        category="attendance",
        link=f"/hr/attendance?month={cal.month_key_of(row.day)}")
    return CorrectionOut.from_model(row)


# --- Export ----------------------------------------------------------------------


@router.get("/export", dependencies=[Depends(require_permission(
    VIEW_ATTENDANCE, EXPORT_DATA))])
async def export_register(month: Optional[str] = Query(default=None),
                          actor: User = Depends(get_active_user)) -> Response:
    """The monthly attendance register (owner H8) — the sheet you print.

    People down the side, days across the top, and the summary columns on the
    right. NO MONEY COLUMN, deliberately (owner 2026-08-20): the manager reads
    `Payable days` off the end of the row and multiplies it themselves.

    Registered BEFORE nothing dynamic — this router has no `/{id}` route for it
    to be shadowed by — but the name is kept explicit because the
    static-before-dynamic scar in this codebase came from exactly this shape.
    """
    now = utcnow()
    hr = (await settings_svc.get_settings()).hr
    key = _month_or_now(month, now)
    employees = await svc.hr_employees()
    first, last = cal.month_bounds(key)
    dates = cal.month_iter(key)

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = f"Attendance {key}"

    head = ["Employee", "Code", "Designation"]
    head += [str(d.day) for d in dates]
    head += ["Present", "Half days", "WFH", "Paid leave", "Unpaid leave",
             "Absent", "Week offs", "Holidays", "Late marks",
             "Late penalty (days)", "Hours worked", "Payable days"]
    ws.append(head)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    # One letter per day, so a month fits on a page. The legend below the grid
    # says what each one is — a grid of codes with no legend is a grid nobody
    # can read without asking.
    letters = {
        AttendanceStatus.PRESENT.value: "P",
        AttendanceStatus.HALF_DAY.value: "H",
        AttendanceStatus.ABSENT.value: "A",
        AttendanceStatus.ON_LEAVE.value: "L",
        AttendanceStatus.LEAVE_UNPAID.value: "LU",
        AttendanceStatus.WEEK_OFF.value: "WO",
        AttendanceStatus.HOLIDAY.value: "HO",
        AttendanceStatus.WFH.value: "W",
        AttendanceStatus.NOT_MARKED.value: "-",
    }
    for user in employees:
        days = await svc.build_month(user, key, hr, now=now)
        s = svc.summarise(days, hr)
        prof = getattr(user, "employee_profile", None)
        row = [user.full_name, user.code, getattr(prof, "designation", "") or ""]
        row += [letters.get(v.status, "-") for v in days]
        row += [s["present"], s["half_days"], s["wfh"], s["on_leave"],
                s["leave_unpaid"], s["absent"], s["week_offs"], s["holidays"],
                s["late_marks"], s["late_penalty_days"], s["worked_hours"],
                s["payable_days"]]
        ws.append(row)

    ws.append([])
    ws.append(["Legend", "P Present", "H Half day", "A Absent",
               "L Paid leave", "LU Unpaid leave", "WO Week off",
               "HO Holiday", "W Work from home", "- Not marked"])
    ws.append([])
    ws.append(["Payable days = calendar days − absent − unpaid leave − half-day "
               "shortfall − late penalty. Salary is worked out from this by "
               "hand."])

    ws.freeze_panes = "D2"
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 18

    buf = io.BytesIO()
    wb.save(buf)
    await log_action(
        AuditAction.REPORT_EXPORTED, actor_id=str(actor.id),
        actor_name=actor.full_name, actor_role=actor.account_type.value,
        entity_type="attendance", summary=f"Exported the {key} attendance "
                                          f"register")
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument."
                   "spreadsheetml.sheet",
        headers={"Content-Disposition":
                 f'attachment; filename="attendance-{key}.xlsx"'})
